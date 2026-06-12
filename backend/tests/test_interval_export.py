"""Per-interval TMC sheet in the v3 Excel export (Phase 4 full-study work).

Synthetic intersection-day: one camera, two legs, a video row (wall-clock
anchor), events across two 15-minute bins separated by a gap — verifying
binning, cardinal grouping, gap omission, and hard-coded integer cells.
"""
import json
import tempfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app import app
from backend.database import get_connection
from backend.services.v3_aggregator import aggregate_intersection_day
from backend.services.v3_excel_export import export_intersection_day_xlsx

client = TestClient(app)


@pytest.fixture()
def interval_project():
    r = client.post("/api/projects", json={"name": "interval-test"})
    pid = r.json()["project_id"]
    conn = get_connection(pid)
    with conn:
        cur = conn.execute(
            "INSERT INTO intersections (name, date, sort_order, leg_count, created_at) "
            "VALUES ('Iv Int', '2026-06-12', 0, 2, '2026-06-12')")
        iid = cur.lastrowid
        cur = conn.execute(
            "INSERT INTO cameras (intersection_id, label, sort_order, created_at) "
            "VALUES (?, 'cam', 0, '2026-06-12')", (iid,))
        cid = cur.lastrowid
        cur = conn.execute(
            "INSERT INTO videos (camera_id, sort_order, path, filename, fps, width, "
            "height, total_frames, duration_seconds, file_size_bytes, "
            "recording_start_datetime, added_at) "
            "VALUES (?, 0, 'x.mp4', 'x.mp4', 10, 640, 480, 864000, 86400, 1, "
            "'2026-05-12T00:00:02', '2026-06-12')",
            (cid,))
        vid = cur.lastrowid
        legs = {}
        for card, label in (("N", "NB Main"), ("S", "SB Main")):
            cur = conn.execute(
                "INSERT INTO legs (camera_id, label, cardinal_direction, sort_order, "
                "origin_zone, reference_heading) VALUES (?, ?, ?, 0, ?, 0)",
                (cid, label, card, json.dumps([[100, 100]])))
            legs[card] = cur.lastrowid
        # Events: 3 NB-through in the 07:00 bin, 2 SB-left in the 07:00 bin,
        # 4 NB-through in the 07:30 bin (07:15 bin left EMPTY = gap).
        # timestamp_video is footage-seconds from 00:00:02.
        def add(ts, o, d, mv, i):
            conn.execute(
                "INSERT INTO vehicle_events (camera_id, video_id, vehicle_track_id, "
                "origin_leg_id, destination_leg_id, movement, trajectory_data, "
                "trajectory_confidence, vehicle_class, detection_confidence, "
                "timestamp_video, frame_number) "
                "VALUES (?, ?, ?, ?, ?, ?, '[]', 1.0, 'car', 0.9, ?, ?)",
                (cid, vid, i, o, d, mv, ts, int(ts * 10)))
        i = 0
        for k in range(3):
            add(25200 + 60 * k, legs["N"], legs["S"], "through", i); i += 1
        for k in range(2):
            add(25260 + 60 * k, legs["S"], legs["N"], "left", i); i += 1
        for k in range(4):
            add(27000 + 30 * k, legs["N"], legs["S"], "through", i); i += 1
    conn.close()
    yield pid, iid
    client.delete(f"/api/projects/{pid}")


class TestIntervalAggregation:
    def test_bins_and_gap(self, interval_project):
        pid, iid = interval_project
        agg = aggregate_intersection_day(pid, iid)
        assert agg["interval_minutes"] == 15
        labels = [iv["label"] for iv in agg["intervals"]]
        # 25200s after 00:00:02 = 07:00:02 -> bin 07:00; 27000s -> 07:30.
        assert labels == ["07:00", "07:30"]          # the 07:15 gap is omitted
        first = agg["intervals"][0]
        assert first["counts"]["N"]["through"] == 3
        assert first["counts"]["S"]["left"] == 2
        assert first["total"] == 5
        assert agg["intervals"][1]["counts"]["N"]["through"] == 4

    def test_excel_interval_sheet(self, interval_project):
        pid, iid = interval_project
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out = Path(tmpdir) / "study.xlsx"
            export_intersection_day_xlsx(pid, iid, out)
            wb = openpyxl.load_workbook(str(out))
            try:
                assert "15-min Intervals" in wb.sheetnames
                ws = wb["15-min Intervals"]
                # Header: approach group on row 1, movements on row 2.
                assert ws.cell(row=1, column=2).value == "NB approach"
                assert ws.cell(row=2, column=2).value == "Thru"
                # Row 3 = 07:00 bin: NB thru 3 ... SB left 2; total 5.
                assert ws.cell(row=3, column=1).value == "07:00"
                assert ws.cell(row=3, column=2).value == 3
                assert ws.cell(row=3, column=7).value == 2   # SB group, Left col
                assert ws.cell(row=3, column=10).value == 5  # row total
                # Row 4 = 07:30 bin; row 5 = bold grand totals.
                assert ws.cell(row=4, column=1).value == "07:30"
                assert ws.cell(row=4, column=2).value == 4
                assert ws.cell(row=5, column=1).value == "Total"
                assert ws.cell(row=5, column=2).value == 7
                assert ws.cell(row=5, column=10).value == 9
                # Hard-coded integers, no formulas (CLAUDE.md constraint).
                for row in ws.iter_rows():
                    for cell in row:
                        assert not (isinstance(cell.value, str)
                                    and cell.value.startswith("="))
            finally:
                wb.close()
