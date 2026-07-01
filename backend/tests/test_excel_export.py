"""Tests for generate_tmc_excel and the export API endpoints."""
import json
import tempfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app import app
from backend.database import get_connection, set_project_info
from backend.services.excel_export import generate_tmc_excel

client = TestClient(app)


def _create_project(name: str = "export-test") -> str:
    r = client.post("/api/projects", json={"name": name})
    assert r.status_code == 200
    return r.json()["project_id"]


def _delete_project(pid: str) -> None:
    client.delete(f"/api/projects/{pid}")


def _seed_data(pid: str) -> None:
    """Insert two legs and a handful of vehicle events for testing."""
    conn = get_connection(pid)
    try:
        conn.execute(
            "INSERT INTO legs (label, cardinal_direction, sort_order, origin_zone, reference_heading) "
            "VALUES (?, ?, ?, ?, ?)",
            ("North", "N", 0, json.dumps([[0.0, 0.0], [100.0, 0.0]]), 0.0),
        )
        conn.execute(
            "INSERT INTO legs (label, cardinal_direction, sort_order, origin_zone, reference_heading) "
            "VALUES (?, ?, ?, ?, ?)",
            ("South", "S", 1, json.dumps([[0.0, 200.0], [100.0, 200.0]]), 180.0),
        )
        north_id = conn.execute("SELECT leg_id FROM legs WHERE label='North'").fetchone()[0]
        south_id = conn.execute("SELECT leg_id FROM legs WHERE label='South'").fetchone()[0]

        events = [
            (1, north_id, "through", "car", None, 0.9, 0.85, 10.0, "2024-01-01 08:00:10", 100, 0),
            (2, north_id, "left",    "car", None, 0.8, 0.80, 20.0, "2024-01-01 08:00:20", 200, 0),
            (3, north_id, "right",   "truck", 9,  0.75, 0.70, 30.0, "2024-01-01 08:00:30", 300, 0),
            (4, south_id, "through", "car", None, 0.9, 0.90, 40.0, "2024-01-01 08:00:40", 400, 0),
            (5, south_id, "u_turn",  "motorcycle", None, 0.6, 0.60, 50.0, "2024-01-01 08:00:50", 500, 0),
        ]
        for e in events:
            conn.execute(
                "INSERT INTO vehicle_events "
                "(vehicle_track_id, origin_leg_id, movement, vehicle_class, fhwa_class, "
                "detection_confidence, trajectory_confidence, timestamp_video, timestamp_real, "
                "frame_number, trajectory_data, manually_edited) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,'[]',?)",
                e,
            )
        conn.commit()
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# Unit tests for generate_tmc_excel


def test_generate_tmc_excel_creates_file():
    pid = _create_project("excel-unit")
    _seed_data(pid)
    try:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out = Path(tmpdir) / "test.xlsx"
            result = generate_tmc_excel(pid, out)
            assert result == out
            assert out.exists()
    finally:
        _delete_project(pid)


def test_excel_sheet_names():
    pid = _create_project("excel-sheets")
    _seed_data(pid)
    try:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out = Path(tmpdir) / "tmc.xlsx"
            generate_tmc_excel(pid, out)
            wb = openpyxl.load_workbook(str(out))
            try:
                assert wb.sheetnames == ["TMC Summary", "Time Series", "TMV Data", "Raw Events"]
            finally:
                wb.close()
    finally:
        _delete_project(pid)


def test_excel_tmc_values_are_integers():
    pid = _create_project("excel-int")
    _seed_data(pid)
    try:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out = Path(tmpdir) / "tmc.xlsx"
            generate_tmc_excel(pid, out)
            wb = openpyxl.load_workbook(str(out))
            try:
                ws = wb["TMC Summary"]
                # Find rows after the header (row with "Leg")
                header_row = None
                for row in ws.iter_rows():
                    if row[0].value == "Leg":
                        header_row = row[0].row
                        break
                assert header_row is not None, "Header row not found"
                # Check all numeric cells in data rows are int, not formula
                for row in ws.iter_rows(min_row=header_row + 1):
                    for cell in row[1:]:  # skip the label column
                        if cell.value is not None:
                            assert isinstance(cell.value, (int, float)), (
                                f"Cell {cell.coordinate} has non-numeric value {cell.value!r}"
                            )
                            assert not str(cell.value).startswith("="), (
                                f"Cell {cell.coordinate} contains a formula"
                            )
            finally:
                wb.close()
    finally:
        _delete_project(pid)


def test_excel_tmc_counts_correct():
    pid = _create_project("excel-counts")
    _seed_data(pid)
    try:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out = Path(tmpdir) / "tmc.xlsx"
            generate_tmc_excel(pid, out)
            wb = openpyxl.load_workbook(str(out))
            try:
                ws = wb["TMC Summary"]

                # Collect data rows between "Leg" header and "Total" footer
                header_row = None
                for row in ws.iter_rows():
                    if row[0].value == "Leg":
                        header_row = row[0].row
                        break
                assert header_row is not None

                north_row = south_row = None
                for row in ws.iter_rows(min_row=header_row + 1):
                    if row[0].value == "North":
                        north_row = [c.value for c in row]
                    elif row[0].value == "South":
                        south_row = [c.value for c in row]

                # North: through=1, left=1, right=1, u_turn=0, total=3
                assert north_row[1] == 1  # through
                assert north_row[2] == 1  # left
                assert north_row[3] == 1  # right
                assert north_row[4] == 0  # u_turn
                assert north_row[5] == 3  # total

                # South: through=1, left=0, right=0, u_turn=1, total=2
                assert south_row[1] == 1
                assert south_row[2] == 0
                assert south_row[3] == 0
                assert south_row[4] == 1
                assert south_row[5] == 2
            finally:
                wb.close()
    finally:
        _delete_project(pid)


def test_excel_tmv_data_sheet_class_aware():
    """The Miovision-format TMV Data sheet: Interval|Approach|Movement|Class|Volume,
    with the truck (fhwa 9) bucketed Articulated and a bound-direction approach."""
    pid = _create_project("excel-tmv")
    _seed_data(pid)
    try:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out = Path(tmpdir) / "tmc.xlsx"
            generate_tmc_excel(pid, out)
            wb = openpyxl.load_workbook(str(out))
            try:
                ws = wb["TMV Data"]
                rows = list(ws.iter_rows(values_only=True))
                assert rows[0] == ("Interval", "Approach", "Movement", "Class", "Volume")
                body = rows[1:]
                classes = {r[3] for r in body}
                assert "Lights" in classes and "Articulated Trucks" in classes
                # the fhwa=9 truck was a North-leg (position N -> Southbound) right turn
                artic = [r for r in body if r[3] == "Articulated Trucks"]
                assert artic and artic[0][1] == "Southbound" and artic[0][2] == "R"
                assert sum(r[4] for r in body) == 5     # all 5 events represented
            finally:
                wb.close()
    finally:
        _delete_project(pid)


def test_excel_no_events_still_creates_file():
    pid = _create_project("excel-empty")
    try:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
            out = Path(tmpdir) / "empty.xlsx"
            result = generate_tmc_excel(pid, out)
            assert result.exists()
            wb = openpyxl.load_workbook(str(out))
            try:
                assert "TMC Summary" in wb.sheetnames
            finally:
                wb.close()
    finally:
        _delete_project(pid)


# ---------------------------------------------------------------------------
# API endpoint tests


def test_export_preview_endpoint():
    pid = _create_project("export-api")
    _seed_data(pid)
    try:
        r = client.get(f"/api/projects/{pid}/export/preview")
        assert r.status_code == 200
        data = r.json()
        assert "tmc_matrix" in data
        assert "total_vehicles" in data
        assert data["total_vehicles"] == 5
        assert data["leg_count"] == 2
        # L/M/A class summary (Miovision parity): 4 unclassified->Lights, 1 fhwa9->Articulated
        assert data["class_summary"] == {"Lights": 4, "Mediums": 0, "Articulated Trucks": 1}
    finally:
        _delete_project(pid)


def test_export_download_endpoint():
    pid = _create_project("export-dl")
    _seed_data(pid)
    try:
        r = client.get(f"/api/projects/{pid}/export/download")
        assert r.status_code == 200
        assert "spreadsheet" in r.headers["content-type"]
        assert len(r.content) > 0
    finally:
        _delete_project(pid)


def test_export_preview_empty_project():
    pid = _create_project("export-empty")
    try:
        r = client.get(f"/api/projects/{pid}/export/preview")
        assert r.status_code == 200
        data = r.json()
        assert data["total_vehicles"] == 0
        assert data["tmc_matrix"] == []
    finally:
        _delete_project(pid)
