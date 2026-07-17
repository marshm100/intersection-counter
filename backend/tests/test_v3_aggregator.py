"""v3 Phase 5 — aggregator + Excel export tests.

Bypasses pipeline + YOLO by writing synthetic vehicle_events rows directly
into the DB. Exercises dedup, aggregation, and the .xlsx exporter shape.
"""

import io
import json
import os
import shutil
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend import database
from backend.app import app
from backend.services.v3_aggregator import aggregate_intersection_day
from backend.services.v3_excel_export import export_intersection_day_xlsx


client = TestClient(app)


def _make_video(dir_path: str, name: str) -> str:
    p = os.path.join(dir_path, name)
    w = cv2.VideoWriter(p, cv2.VideoWriter_fourcc(*"mp4v"), 30.0, (320, 240))
    for _ in range(90):
        w.write(np.zeros((240, 320, 3), dtype=np.uint8))
    w.release()
    return p


def _insert_event(
    pid: str, *, video_id: int, camera_id: int, trim_id: int | None,
    leg_id: int, movement: str, timestamp_video: float = 1.0,
    detection_confidence: float = 0.9, rejected: int = 0,
) -> int:
    conn = database.get_connection(pid)
    try:
        cur = conn.execute(
            """INSERT INTO vehicle_events
               (video_id, camera_id, trim_id, vehicle_track_id,
                origin_leg_id, movement, trajectory_data,
                trajectory_confidence, vehicle_class, detection_confidence,
                timestamp_video, frame_number, rejected)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (video_id, camera_id, trim_id, 1, leg_id, movement, "[]",
             0.9, "car", detection_confidence,
             timestamp_video, int(timestamp_video * 30), rejected),
        )
        conn.commit()
        return int(cur.lastrowid)
    finally:
        conn.close()


def _insert_leg(pid: str, *, camera_id: int, label: str, cardinal: str,
                sort_order: int = 0) -> int:
    conn = database.get_connection(pid)
    try:
        cur = conn.execute(
            """INSERT INTO legs
               (camera_id, label, cardinal_direction, sort_order,
                origin_zone, reference_heading)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (camera_id, label, cardinal, sort_order, "[]", 0.0),
        )
        conn.commit()
        return int(cur.lastrowid)
    finally:
        conn.close()


@pytest.fixture()
def two_cam_intersection():
    """Project with 2 cameras at one intersection-day, each with one leg."""
    pid = client.post("/api/projects", json={"name": "v3-agg"}).json()["project_id"]
    tmpdir = tempfile.mkdtemp(prefix="v3_agg_")
    try:
        paths = [
            _make_video(tmpdir, "Cam1_01_20260514_080000 Main.mp4"),
            _make_video(tmpdir, "Cam2_01_20260514_080000 Main.mp4"),
        ]
        client.post(f"/api/projects/{pid}/videos/bulk", json={"paths": paths})
        body = client.post(f"/api/projects/{pid}/videos/save-labels").json()
        iid = body["intersections"][0]["intersection_id"]
        cams = client.get(f"/api/projects/{pid}/intersections/{iid}/cameras").json()
        cam_a, cam_b = cams[0]["camera_id"], cams[1]["camera_id"]
        leg_a = _insert_leg(pid, camera_id=cam_a, label="North", cardinal="N")
        leg_b = _insert_leg(pid, camera_id=cam_b, label="North", cardinal="N")
        vids = client.get(f"/api/projects/{pid}/videos").json()
        vid_a = next(v["video_id"] for v in vids if v["camera_id"] == cam_a)
        vid_b = next(v["video_id"] for v in vids if v["camera_id"] == cam_b)
        yield {"pid": pid, "iid": iid,
               "cam_a": cam_a, "cam_b": cam_b,
               "leg_a": leg_a, "leg_b": leg_b,
               "vid_a": vid_a, "vid_b": vid_b}
    finally:
        client.delete(f"/api/projects/{pid}")
        shutil.rmtree(tmpdir, ignore_errors=True)


class TestAggregation:
    def test_no_events_returns_zeroed_matrix(self, two_cam_intersection):
        ctx = two_cam_intersection
        agg = aggregate_intersection_day(ctx["pid"], ctx["iid"])
        # One row per configured leg label even with no events, so the live
        # mid-run dashboard shows a complete TMC table from t=0.
        assert agg["totals"]["vehicles"] == 0
        assert [r["leg_label"] for r in agg["tmc_matrix"]] == ["North"]
        assert all(r["total"] == 0 for r in agg["tmc_matrix"])

    def test_singleton_event_counted(self, two_cam_intersection):
        ctx = two_cam_intersection
        _insert_event(
            ctx["pid"],
            video_id=ctx["vid_a"], camera_id=ctx["cam_a"], trim_id=None,
            leg_id=ctx["leg_a"], movement="through", timestamp_video=1.0,
        )
        agg = aggregate_intersection_day(ctx["pid"], ctx["iid"])
        assert agg["totals"]["vehicles"] == 1
        assert agg["tmc_matrix"][0]["leg_label"] == "North"
        assert agg["tmc_matrix"][0]["through"] == 1

    def test_rejected_events_excluded(self, two_cam_intersection):
        ctx = two_cam_intersection
        _insert_event(
            ctx["pid"], video_id=ctx["vid_a"], camera_id=ctx["cam_a"],
            trim_id=None, leg_id=ctx["leg_a"], movement="through",
            timestamp_video=1.0, rejected=1,
        )
        agg = aggregate_intersection_day(ctx["pid"], ctx["iid"])
        assert agg["totals"]["vehicles"] == 0

    def test_parallel_overlap_deduplicates(self, two_cam_intersection):
        ctx = two_cam_intersection
        # Cam A and Cam B both see a left turn within 1 second of each other,
        # during their parallel-overlap window. Should count as 1.
        _insert_event(
            ctx["pid"], video_id=ctx["vid_a"], camera_id=ctx["cam_a"],
            trim_id=None, leg_id=ctx["leg_a"], movement="left",
            timestamp_video=1.0, detection_confidence=0.85,
        )
        _insert_event(
            ctx["pid"], video_id=ctx["vid_b"], camera_id=ctx["cam_b"],
            trim_id=None, leg_id=ctx["leg_b"], movement="left",
            timestamp_video=1.5, detection_confidence=0.95,
        )
        agg = aggregate_intersection_day(ctx["pid"], ctx["iid"])
        assert agg["totals"]["vehicles"] == 1
        assert agg["dedup_summary"]["merged"] == 1

    def test_per_camera_breakdown_shows_raw_counts(self, two_cam_intersection):
        """Per-camera sheet shows raw counts (not deduped) for QA."""
        ctx = two_cam_intersection
        _insert_event(ctx["pid"], video_id=ctx["vid_a"], camera_id=ctx["cam_a"],
                      trim_id=None, leg_id=ctx["leg_a"], movement="left",
                      timestamp_video=1.0, detection_confidence=0.85)
        _insert_event(ctx["pid"], video_id=ctx["vid_b"], camera_id=ctx["cam_b"],
                      trim_id=None, leg_id=ctx["leg_b"], movement="left",
                      timestamp_video=1.5, detection_confidence=0.95)
        agg = aggregate_intersection_day(ctx["pid"], ctx["iid"])
        # Each camera shows 1 raw event in the breakdown sheet
        for cam_block in agg["per_camera_breakdown"]:
            assert cam_block["raw_event_count"] == 1


class TestExcelExport:
    def test_exports_three_sheets(self, two_cam_intersection, tmp_path):
        ctx = two_cam_intersection
        _insert_event(ctx["pid"], video_id=ctx["vid_a"], camera_id=ctx["cam_a"],
                      trim_id=None, leg_id=ctx["leg_a"], movement="through",
                      timestamp_video=1.0)
        out = tmp_path / "out.xlsx"
        export_intersection_day_xlsx(ctx["pid"], ctx["iid"], out)
        assert out.exists()
        wb = openpyxl.load_workbook(out)
        assert "TMC Summary" in wb.sheetnames
        assert "Per-Camera Breakdown" in wb.sheetnames
        assert "Dedup Audit" in wb.sheetnames

    def test_summary_sheet_values(self, two_cam_intersection, tmp_path):
        ctx = two_cam_intersection
        _insert_event(ctx["pid"], video_id=ctx["vid_a"], camera_id=ctx["cam_a"],
                      trim_id=None, leg_id=ctx["leg_a"], movement="through",
                      timestamp_video=1.0)
        _insert_event(ctx["pid"], video_id=ctx["vid_a"], camera_id=ctx["cam_a"],
                      trim_id=None, leg_id=ctx["leg_a"], movement="left",
                      timestamp_video=1.5)
        out = tmp_path / "out.xlsx"
        export_intersection_day_xlsx(ctx["pid"], ctx["iid"], out)
        wb = openpyxl.load_workbook(out)
        ws = wb["TMC Summary"]
        # Header row at row 6 (after 4 metadata + 1 blank)
        # First data row at 7 — should have leg label North, through=1, left=1
        row = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == "North")
        assert row[1] == 1   # Through
        assert row[2] == 1   # Left

    def test_export_endpoint_returns_xlsx(self, two_cam_intersection):
        ctx = two_cam_intersection
        _insert_event(ctx["pid"], video_id=ctx["vid_a"], camera_id=ctx["cam_a"],
                      trim_id=None, leg_id=ctx["leg_a"], movement="through",
                      timestamp_video=1.0)
        r = client.get(
            f"/api/projects/{ctx['pid']}/intersections/{ctx['iid']}/export/xlsx"
        )
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        # Body should be a real xlsx (zip starts with PK)
        assert r.content[:2] == b"PK"


class TestExportLoaderV3:
    """Stage 2a (plan_deliverables_E_2026-07-16): the shared v3 export frame —
    deduped, rejected-excluded, per-event wall-clock, no epoch fallback."""

    def test_dedup_and_leg_merge(self, two_cam_intersection):
        s = two_cam_intersection
        from backend.services.excel_export import _load_export_data_v3
        # the same physical vehicle seen by both cameras in the overlap
        _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                      trim_id=None, leg_id=s["leg_a"], movement="through",
                      timestamp_video=1.0)
        _insert_event(s["pid"], video_id=s["vid_b"], camera_id=s["cam_b"],
                      trim_id=None, leg_id=s["leg_b"], movement="through",
                      timestamp_video=1.0)
        d = _load_export_data_v3(s["pid"], s["iid"])
        # one merged "North" leg row, not one per camera
        assert [lg[1] for lg in d["legs"]].count("North") == 1
        north = next(v for v in d["tmc"].values() if v["label"] == "North")
        assert north["through"] == 1          # deduped to one vehicle
        assert len(d["events"]) == 1

    def test_rejected_excluded(self, two_cam_intersection):
        s = two_cam_intersection
        from backend.services.excel_export import _load_export_data_v3
        _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                      trim_id=None, leg_id=s["leg_a"], movement="left",
                      timestamp_video=1.0, rejected=1)
        d = _load_export_data_v3(s["pid"], s["iid"])
        assert len(d["events"]) == 0
        assert all(v["total"] == 0 for v in d["tmc"].values())

    def test_no_epoch_intervals(self, two_cam_intersection):
        s = two_cam_intersection
        from backend.services.excel_export import _load_export_data_v3
        _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                      trim_id=None, leg_id=s["leg_a"], movement="through",
                      timestamp_video=61.0)
        d = _load_export_data_v3(s["pid"], s["iid"])
        assert d["n_unstamped"] == 0
        assert d["tmv"], "event must reach the tmv frame"
        for (interval, _a, _m, _c) in d["tmv"]:
            assert interval.startswith("2026-05-14 08:"), interval

    def test_unstamped_counted_not_binned(self, two_cam_intersection):
        s = two_cam_intersection
        from backend.services.excel_export import _load_export_data_v3
        # no video row -> no wall-clock: counted in tmc, EXCLUDED from tmv
        _insert_event(s["pid"], video_id=None, camera_id=s["cam_a"],
                      trim_id=None, leg_id=s["leg_a"], movement="through",
                      timestamp_video=5.0)
        d = _load_export_data_v3(s["pid"], s["iid"])
        assert d["n_unstamped"] == 1
        assert len(d["events"]) == 1
        assert not d["tmv"]
        north = next(v for v in d["tmc"].values() if v["label"] == "North")
        assert north["through"] == 1

    def test_generate_excel_v3_frame(self, two_cam_intersection, tmp_path):
        s = two_cam_intersection
        from backend.services.excel_export import generate_tmc_excel
        _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                      trim_id=None, leg_id=s["leg_a"], movement="through",
                      timestamp_video=1.0)
        out = generate_tmc_excel(s["pid"], tmp_path / "v3.xlsx",
                                 intersection_id=s["iid"])
        wb = openpyxl.load_workbook(out)
        assert "TMC Summary" in wb.sheetnames


class TestMiovisionWorkbook:
    """Stage 2b: the Miovision-parity workbook structure pins."""

    def _mk(self, s, tmp_path):
        from backend.services.excel_export import generate_miovision_xlsx
        # AM-peak-hour traffic: through + left on North; NO right all day
        for k in range(8):
            _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                          trim_id=None, leg_id=s["leg_a"],
                          movement="through" if k % 2 else "left",
                          timestamp_video=float(k))
        return generate_miovision_xlsx(s["pid"], tmp_path / "mio.xlsx",
                                       s["iid"])

    def test_sheet_inventory_and_headers(self, two_cam_intersection, tmp_path):
        s = two_cam_intersection
        out = self._mk(s, tmp_path)
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames[0] == "Contents"
        assert wb.sheetnames[1].endswith(") Summary")
        assert wb.sheetnames[2:] == ["TMV Table", "TMV Data", "Raw Events"]
        for name in wb.sheetnames[:4]:
            ws = wb[name]
            assert ws["B1"].value == "Study Name"
            assert ws["B2"].value == "Start Date"
            assert ws["B3"].value == "End Date"
            assert ws["B4"].value == "Site Code"

    def test_geometry_aware_letters(self, two_cam_intersection, tmp_path):
        s = two_cam_intersection
        out = self._mk(s, tmp_path)
        wb = openpyxl.load_workbook(out)
        ws = wb[wb.sheetnames[1]]
        # letters row: observed movements (L, T) + U always; R absent
        letters = [ws.cell(9, c).value for c in range(4, 10)
                   if ws.cell(9, c).value]
        assert letters[:3] == ["L", "T", "U"]
        assert "R" not in letters
        assert letters[3:5] == ["I", "O"]

    def test_no_formulas_anywhere(self, two_cam_intersection, tmp_path):
        s = two_cam_intersection
        out = self._mk(s, tmp_path)
        wb = openpyxl.load_workbook(out)
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    assert not (isinstance(cell.value, str)
                                and cell.value.startswith("=")), \
                        f"formula at {name}!{cell.coordinate}"

    def test_tmv_data_long_format(self, two_cam_intersection, tmp_path):
        s = two_cam_intersection
        out = self._mk(s, tmp_path)
        wb = openpyxl.load_workbook(out)
        ws = wb["TMV Data"]
        hdr = [ws.cell(8, c).value for c in range(2, 7)]
        assert hdr == ["Interval", "Approach", "Movement", "Class", "Volume"]
        # deduped: 8 inserted events on one camera -> 2 (approach, movement)
        # aggregate rows in the 08:00 interval
        rows = [[ws.cell(r, c).value for c in range(2, 7)]
                for r in range(9, 12) if ws.cell(r, 2).value]
        assert all(v[0].startswith("2026-05-14 08:00") for v in rows)
        assert sum(v[4] for v in rows) == 8


class TestPdfReportV3:
    """Stage 2c: the rebuilt PDF — per-minute pages, grouped columns,
    letterhead config semantics."""

    def _mk(self, s, tmp_path, monkeypatch=None, letterhead=None):
        from backend.services import pdf_report
        if monkeypatch is not None:
            monkeypatch.setattr(pdf_report, "REPORT_LETTERHEAD", letterhead or {
                "name": "", "address_lines": [], "contact": "", "tagline": ""})
        for k in range(6):
            _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                          trim_id=None, leg_id=s["leg_a"], movement="through",
                          timestamp_video=float(k * 30))
        return pdf_report.generate_report_pdf(s["pid"], tmp_path / "r.pdf",
                                              intersection_id=s["iid"])

    def test_pages_and_layout(self, two_cam_intersection, tmp_path, monkeypatch):
        from pypdf import PdfReader
        s = two_cam_intersection
        out = self._mk(s, tmp_path, monkeypatch)
        r = PdfReader(str(out))
        # per-minute page + 15-min page + summary
        assert len(r.pages) == 3
        t0 = r.pages[0].extract_text() or ""
        assert "Turning Movement Data" in t0
        assert "App. Total" in t0 and "Int." in t0
        assert "North" in t0                       # the road label
        assert "Page No: 1" in t0
        t1 = r.pages[1].extract_text() or ""
        assert "15 Minute Intervals" in t1
        assert "Page No: 2" in t1

    def test_letterhead_omitted_when_unset(self, two_cam_intersection,
                                           tmp_path, monkeypatch):
        from pypdf import PdfReader
        s = two_cam_intersection
        out = self._mk(s, tmp_path, monkeypatch)
        t = PdfReader(str(out)).pages[0].extract_text() or ""
        assert "Count Name:" in t

    def test_letterhead_rendered_when_set(self, two_cam_intersection,
                                          tmp_path, monkeypatch):
        from pypdf import PdfReader
        s = two_cam_intersection
        out = self._mk(s, tmp_path, monkeypatch, letterhead={
            "name": "Acme Traffic", "address_lines": ["1 Main St"],
            "contact": "555-0100", "tagline": "Counting responsibly."})
        t = PdfReader(str(out)).pages[0].extract_text() or ""
        assert "Acme Traffic" in t and "1 Main St" in t


class TestIntersectionExportEndpoints:
    """Stage 3: per-intersection-day deliverable endpoints + gate."""

    def test_gate_scoped_to_intersection(self, two_cam_intersection):
        s = two_cam_intersection
        r = client.get(f"/api/projects/{s['pid']}/intersections/{s['iid']}"
                       f"/export/gate")
        assert r.status_code == 200
        body = r.json()
        assert body["intersection_id"] == s["iid"]
        assert "blocking" in body and "blocking_reasons" in body

    def test_gate_unknown_intersection_404(self, two_cam_intersection):
        s = two_cam_intersection
        r = client.get(f"/api/projects/{s['pid']}/intersections/99999"
                       f"/export/gate")
        assert r.status_code == 404

    def test_blocked_download_409_then_override(self, two_cam_intersection):
        s = two_cam_intersection
        # fresh fixture: no bank, no classification -> gate blocks
        _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                      trim_id=None, leg_id=s["leg_a"], movement="through",
                      timestamp_video=1.0)
        base = (f"/api/projects/{s['pid']}/intersections/{s['iid']}"
                f"/export/tmc.xlsx")
        r = client.get(base)
        assert r.status_code == 409
        assert "blocking_reasons" in r.json()["detail"]
        r2 = client.get(base + "?override=true")
        assert r2.status_code == 200
        assert r2.headers["content-type"].startswith(
            "application/vnd.openxmlformats")

    def test_pdf_download_with_override(self, two_cam_intersection):
        s = two_cam_intersection
        _insert_event(s["pid"], video_id=s["vid_a"], camera_id=s["cam_a"],
                      trim_id=None, leg_id=s["leg_a"], movement="through",
                      timestamp_video=1.0)
        r = client.get(f"/api/projects/{s['pid']}/intersections/{s['iid']}"
                       f"/export/report.pdf?override=true")
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/pdf"


class TestLabelerEndpoints:
    """Fine-tune labeler backend (plan_detector_finetune): manifest ordering,
    label save round-trip, traversal guard."""

    @pytest.fixture()
    def dataset(self, tmp_path, monkeypatch):
        from backend.routers import labeler as mod
        root = tmp_path / "ftds"
        for sub in ("images/train", "images/val", "labels/train", "labels/val"):
            (root / sub).mkdir(parents=True)
        import numpy as np, cv2 as _cv2
        _cv2.imwrite(str(root / "images/train/a.jpg"),
                     np.zeros((48, 64, 3), dtype=np.uint8))
        _cv2.imwrite(str(root / "images/train/b.jpg"),
                     np.zeros((48, 64, 3), dtype=np.uint8))
        (root / "labels/train/a.txt").write_text("0 0.5 0.5 0.10 0.2\n")
        (root / "labels/train/b.txt").write_text("0 0.5 0.5 0.40 0.2\n")
        monkeypatch.setattr(mod, "_DATA_ROOT", tmp_path.resolve())
        return "ftds"

    def test_manifest_truck_first_order(self, dataset):
        m = client.get(f"/api/labeler/{dataset}/manifest").json()
        assert [it["name"] for it in m["items"]] == ["b.jpg", "a.jpg"]  # widest first
        assert m["total"] == 2 and m["reviewed_count"] == 0

    def test_save_roundtrip_and_reviewed(self, dataset):
        r = client.post(f"/api/labeler/{dataset}/labels", json={
            "split": "train", "name": "a.jpg",
            "boxes": [[1, 0.5, 0.5, 0.1, 0.2], [0, 0.2, 0.2, 0.05, 0.05]]})
        assert r.status_code == 200 and r.json()["boxes"] == 2
        m = client.get(f"/api/labeler/{dataset}/manifest").json()
        a = next(it for it in m["items"] if it["name"] == "a.jpg")
        assert a["reviewed"] is True
        assert [b[0] for b in a["boxes"]] == [1, 0]

    def test_bad_class_rejected(self, dataset):
        r = client.post(f"/api/labeler/{dataset}/labels", json={
            "split": "train", "name": "a.jpg", "boxes": [[7, .5, .5, .1, .1]]})
        assert r.status_code == 422

    def test_medium_class_accepted(self, dataset):
        # round-v2 taxonomy: 3 = medium (bus / non-long single-unit truck)
        r = client.post(f"/api/labeler/{dataset}/labels", json={
            "split": "train", "name": "b.jpg", "boxes": [[3, .5, .5, .1, .1]]})
        assert r.status_code == 200 and r.json()["boxes"] == 1

    def test_traversal_guarded(self, dataset):
        assert client.get("/api/labeler/..%2f..%2fetc/manifest").status_code == 404
