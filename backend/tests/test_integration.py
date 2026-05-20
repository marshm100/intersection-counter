"""End-to-end integration tests.

Drives a synthetic vehicle through the full processing chain:
  synthetic video → pipeline (mocked detector/tracker) → DB → Excel export

No GPU, no real YOLO weights, no real video required.
"""

import json
import sqlite3
import tempfile
from pathlib import Path

import cv2
import numpy as np
import openpyxl
import pytest
from fastapi.testclient import TestClient
from unittest.mock import MagicMock

from backend.app import app
from backend.database import get_connection, get_db_path, SCHEMA
from backend.services.excel_export import generate_tmc_excel
from backend.services.pipeline import ProcessingPipeline

client = TestClient(app)


# ---------------------------------------------------------------------------
# Helpers shared by both tests
# ---------------------------------------------------------------------------

def _create_project(name: str = "integration-test") -> str:
    r = client.post("/api/projects", json={"name": name})
    assert r.status_code == 200
    return r.json()["project_id"]


def _delete_project(pid: str) -> None:
    client.delete(f"/api/projects/{pid}")


def _make_detection(x: float, y: float, w: float = 100, h: float = 60,
                    class_id: int = 2, confidence: float = 0.9) -> dict:
    """Build a detection dict at center (x, y)."""
    x1, y1, x2, y2 = x - w / 2, y - h / 2, x + w / 2, y + h / 2
    return {
        "bbox": [x1, y1, x2, y2],
        "center": [float(x), float(y)],
        "class_id": class_id,
        "class_name": "car",
        "confidence": confidence,
        "bbox_width": float(w),
        "bbox_height": float(h),
        "bbox_area": float(w * h),
        "is_vehicle": class_id in (2, 3, 5, 7),
    }


def _get_vehicle_events(db_path: str) -> list[dict]:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    rows = conn.execute("SELECT * FROM vehicle_events").fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ---------------------------------------------------------------------------
# Test class
# ---------------------------------------------------------------------------

class TestEndToEndIntegration:

    def test_pipeline_to_excel(self):
        """Full stack: synthetic video → pipeline → DB event → Excel export."""
        pid = _create_project("e2e-pipeline")
        try:
            db_path = str(get_db_path(pid))

            # Insert 4 legs (N/S/E/W). derive_movement's rank-based logic
            # only returns "through" when there are ≥3 other legs at the
            # intersection; with fewer, it collapses to u_turn or a turn.
            # Origin leg = Northbound (vehicle moves south).
            conn = get_connection(pid)
            try:
                leg_seed = [
                    ("Northbound", "NB", 0, [[0, 240], [640, 240]], 180.0),
                    ("Southbound", "SB", 1, [[0, 440], [640, 440]],   0.0),
                    ("Eastbound",  "EB", 2, [[440, 0], [440, 480]], 270.0),
                    ("Westbound",  "WB", 3, [[200, 0], [200, 480]],  90.0),
                ]
                for label, cd, so, oz, rh in leg_seed:
                    conn.execute(
                        "INSERT INTO legs "
                        "(label, cardinal_direction, sort_order, origin_zone, reference_heading) "
                        "VALUES (?, ?, ?, ?, ?)",
                        (label, cd, so, json.dumps(oz), rh),
                    )
                conn.commit()
                leg_id = conn.execute(
                    "SELECT leg_id FROM legs WHERE label='Northbound'"
                ).fetchone()[0]
            finally:
                conn.close()

            with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
                # Create a 100-frame synthetic video (640×480 black frames)
                video_path = str(Path(tmpdir) / "test.mp4")
                fourcc = cv2.VideoWriter_fourcc(*"mp4v")
                writer = cv2.VideoWriter(video_path, fourcc, 30.0, (640, 480))
                for _ in range(100):
                    writer.write(np.zeros((480, 640, 3), dtype=np.uint8))
                writer.release()

                conn = get_connection(pid)
                conn.row_factory = sqlite3.Row
                try:
                    legs = [
                        dict(r) for r in conn.execute(
                            "SELECT leg_id, label, cardinal_direction, "
                            "origin_zone, reference_heading FROM legs "
                            "ORDER BY sort_order"
                        ).fetchall()
                    ]
                finally:
                    conn.close()
                for leg in legs:
                    leg["origin_zone"] = json.loads(leg["origin_zone"])

                pipeline = ProcessingPipeline(
                    project_id=pid,
                    db_path=db_path,
                    video_path=video_path,
                    legs=legs,
                    fps=30.0,
                )

                # Build per-processed-frame detection list.
                # Vehicle moves south (y increasing) crossing y=240 at processed frame 3.
                # Processed frame N corresponds to video frame N*3 (frame_skip=3).
                #
                # Processed 0 (y=215): above zone, no prev → no crossing check
                # Processed 1 (y=225): above zone, prev=215 → no cross
                # Processed 2 (y=230): above zone, prev=225 → no cross
                # Processed 3 (y=250): crosses y=240 (prev=230) → origin assigned,
                #                      trajectory = [(320,250)]
                # Processed 4–33 (y=265,280,…): 30 more trajectory points, vehicle
                #                  stays present through all 34 processed frames so
                #                  _finalize_all_active() is exercised (not lost_ids).
                #
                # Total trajectory: 31 pts ≥ TRAJECTORY_MIN_POINTS=10
                # Distance ≥ TRAJECTORY_MIN_DISTANCE_PX=50
                # Heading: south ≈ 180° = reference_heading → "through"
                above_ys = [215, 225, 230]
                after_ys = [250] + [265 + 15 * i for i in range(30)]  # frames 3-33
                all_ys = above_ys + after_ys   # 34 entries — vehicle present every frame
                detection_seq = [[_make_detection(320, y)] for y in all_ys]

                detect_idx = [0]

                pipeline._preprocessor = MagicMock()
                pipeline._preprocessor.preprocess = MagicMock(side_effect=lambda f: f)

                pipeline._detector = MagicMock()
                def mock_detect(frame):
                    idx = detect_idx[0]
                    detect_idx[0] += 1
                    return detection_seq[idx] if idx < len(detection_seq) else []
                pipeline._detector.detect = mock_detect

                pipeline._tracker = MagicMock()
                def mock_tracker_update(detections, frame_number):
                    return [{**d, "track_id": 1} for d in detections]
                pipeline._tracker.update = mock_tracker_update
                pipeline._tracker.get_state = MagicMock(return_value=b"")

                pipeline.process_video(frame_skip=3)

                # Verify exactly 1 event with movement="through"
                events = _get_vehicle_events(db_path)
                assert len(events) == 1, f"Expected 1 event, got {len(events)}: {events}"
                assert events[0]["movement"] == "through", (
                    f"Expected 'through', got {events[0]['movement']!r}"
                )
                assert events[0]["origin_leg_id"] == leg_id

                # Seed 2 more events directly (left, right)
                conn = get_connection(pid)
                try:
                    for track_id, movement in ((2, "left"), (3, "right")):
                        conn.execute(
                            "INSERT INTO vehicle_events "
                            "(vehicle_track_id, origin_leg_id, movement, trajectory_data, "
                            "trajectory_confidence, vehicle_class, fhwa_class, "
                            "detection_confidence, timestamp_video, timestamp_real, "
                            "frame_number, manually_edited) "
                            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                            (track_id, leg_id, movement, "[]",
                             0.8, "car", None, 0.85, 10.0, None, 100, 0),
                        )
                    conn.commit()
                finally:
                    conn.close()

                # Generate Excel and validate
                output_path = Path(tmpdir) / "tmc_output.xlsx"
                generate_tmc_excel(pid, output_path)

                assert output_path.exists(), "Excel file was not created"

                wb = openpyxl.load_workbook(str(output_path))
                try:
                    assert "TMC Summary" in wb.sheetnames, (
                        f"'TMC Summary' sheet missing; sheets={wb.sheetnames}"
                    )

                    ws_summary = wb["TMC Summary"]

                    # Locate header row (row with first cell value "Leg")
                    header_row = None
                    for row in ws_summary.iter_rows():
                        if row[0].value == "Leg":
                            header_row = row[0].row
                            break
                    assert header_row is not None, "Header row ('Leg') not found"

                    # Collect all data rows after header
                    # Grand total row label is "Total" (matches excel_export.py)
                    data_rows: dict[str, list] = {}
                    grand_total_row = None
                    for row in ws_summary.iter_rows(min_row=header_row + 1):
                        label = row[0].value
                        if label is None:
                            continue
                        vals = [c.value for c in row]
                        label_str = str(label).strip()
                        if label_str == "Total":
                            grand_total_row = vals
                        else:
                            data_rows[label_str] = vals

                    # NB row: Through=1, Left=1, Right=1, U-Turn=0, Total=3
                    nb_row = data_rows.get("Northbound") or data_rows.get("NB")
                    assert nb_row is not None, (
                        f"NB leg row not found; data_rows={list(data_rows.keys())}"
                    )
                    # Columns: [Leg, Through, Left, Right, U-Turn, Total]
                    assert nb_row[1] == 1, f"Through expected 1, got {nb_row[1]}"
                    assert nb_row[2] == 1, f"Left expected 1, got {nb_row[2]}"
                    assert nb_row[3] == 1, f"Right expected 1, got {nb_row[3]}"
                    assert nb_row[4] == 0, f"U-Turn expected 0, got {nb_row[4]}"
                    assert nb_row[5] == 3, f"Total expected 3, got {nb_row[5]}"

                    # All numeric cells in data rows must be int (no formulas, no floats)
                    for row in ws_summary.iter_rows(min_row=header_row + 1):
                        for cell in row[1:]:
                            if cell.value is not None:
                                assert isinstance(cell.value, int), (
                                    f"Cell {cell.coordinate} value {cell.value!r} is not int"
                                )

                    # Grand Total row: Through=1, Left=1, Right=1, U-Turn=0, Total=3
                    assert grand_total_row is not None, "Grand Total row not found"
                    assert grand_total_row[1] == 1, f"GT Through expected 1, got {grand_total_row[1]}"
                    assert grand_total_row[2] == 1, f"GT Left expected 1, got {grand_total_row[2]}"
                    assert grand_total_row[3] == 1, f"GT Right expected 1, got {grand_total_row[3]}"
                    assert grand_total_row[4] == 0, f"GT U-Turn expected 0, got {grand_total_row[4]}"
                    assert grand_total_row[5] == 3, f"GT Total expected 3, got {grand_total_row[5]}"

                    # "Raw Events" sheet has 3 data rows (1 from pipeline + 2 seeded)
                    assert "Raw Events" in wb.sheetnames, "'Raw Events' sheet missing"
                    ws_raw = wb["Raw Events"]
                    data_rows_raw = [
                        row for row in ws_raw.iter_rows(min_row=2, values_only=True)
                        if any(v is not None for v in row)
                    ]
                    assert len(data_rows_raw) == 3, (
                        f"Expected 3 data rows in Raw Events, got {len(data_rows_raw)}"
                    )

                finally:
                    wb.close()

        finally:
            _delete_project(pid)

    def test_export_preview_api(self):
        """HTTP export-preview endpoint returns correct TMC counts for seeded events."""
        pid = _create_project("e2e-preview")
        try:
            conn = get_connection(pid)
            try:
                conn.execute(
                    "INSERT INTO legs "
                    "(label, cardinal_direction, sort_order, origin_zone, reference_heading) "
                    "VALUES (?, ?, ?, ?, ?)",
                    ("North", "NB", 0, json.dumps([[0.0, 240.0], [640.0, 240.0]]), 180.0),
                )
                conn.commit()
                leg_id = conn.execute(
                    "SELECT leg_id FROM legs WHERE label='North'"
                ).fetchone()[0]

                # Seed 2 vehicle events: one "through", one "left"
                for track_id, movement in ((1, "through"), (2, "left")):
                    conn.execute(
                        "INSERT INTO vehicle_events "
                        "(vehicle_track_id, origin_leg_id, movement, trajectory_data, "
                        "trajectory_confidence, vehicle_class, fhwa_class, "
                        "detection_confidence, timestamp_video, timestamp_real, "
                        "frame_number, manually_edited) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                        (track_id, leg_id, movement, "[]",
                         0.9, "car", None, 0.85, 10.0, None, 100, 0),
                    )
                conn.commit()
            finally:
                conn.close()

            r = client.get(f"/api/projects/{pid}/export/preview")
            assert r.status_code == 200, f"Expected 200, got {r.status_code}: {r.text}"

            data = r.json()
            assert "tmc_matrix" in data, f"'tmc_matrix' not in response: {data}"

            matrix = data["tmc_matrix"]
            assert len(matrix) == 1, f"Expected 1 leg row in tmc_matrix, got {len(matrix)}"

            leg_row = matrix[0]
            assert leg_row.get("through") == 1, (
                f"Expected through=1, got {leg_row}"
            )
            assert leg_row.get("left") == 1, (
                f"Expected left=1, got {leg_row}"
            )

        finally:
            _delete_project(pid)
