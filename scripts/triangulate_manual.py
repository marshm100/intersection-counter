"""Validate a corridor camera against its MANUAL hand count (top ground truth).

For the SAME 2026-05-12 footage we can have three counts per intersection:
  - MANUAL    — a hand count (Simeon Lewis), the top ground truth.
  - MIOVISION — the commercial per-minute OD XML (the tool we replace).
  - OURS      — our pipeline's vehicle_events in project.db.

The manual sheets lay out AM+PM peak windows but only a CONTIGUOUS BLOCK of
minutes at the start is actually filled in (cam2: 07:00-07:35; the rest are
blank time labels). So the honest comparison restricts Miovision and OURS to the
exact minutes the human counted, and treats MANUAL as the reference.

Keyed on travel direction (NB/SB/EB/WB) x movement.
  - MANUAL: direction from the approach NAME prefix ('SB Nbeltline Rd' -> SB).
  - MIOVISION: from its approach Name prefix.
  - OURS: each leg's cardinal_direction is its POSITION, so the bound approach is
    the OPPOSITE (south-arm leg cardinal 'S' -> NB) via _CARD_TO_DIR. For cam1
    (generic 'Leg N' labels) this also tests whether our directions are right: a
    clean swap vs MANUAL would show as NB<->SB / EB<->WB mismatches.

Usage:
  py scripts/triangulate_manual.py                 # every camera with a manual file
  py scripts/triangulate_manual.py --camera 1
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from collections import defaultdict
from datetime import datetime, time
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import parse_miovision_xml as MIO

PROJECT = "97a7849a"
_HIST = Path("docs/historic data/Sunnyvale, TX")
MANUAL_BY_CAM = {
    1: _HIST / "NBeltlineRd-NorthwestDr_ManualCounts.xlsx",
    2: _HIST / "NBeltlineRd-ETownEastBlvd_ManualCounts.xlsx",
}

DIRECTIONS = ["NB", "SB", "EB", "WB"]
MOVEMENTS = ["thru", "left", "right", "uturn"]
_MV_FROM_LETTER = {"R": "right", "T": "thru", "L": "left", "U": "uturn"}
# Cardinal is the leg POSITION; the bound approach (what Miovision + the manual
# sheets report) is the OPPOSITE — a south-arm leg carries NB traffic. Mirrors
# backend/services/cardinals.bound_approach.
_CARD_TO_DIR = {"N": "SB", "S": "NB", "E": "WB", "W": "EB",
                "NE": "SWB", "SW": "NEB", "NW": "SEB", "SE": "NWB"}
PerMin = dict[time, dict[tuple, int]]


def parse_manual(path: Path) -> PerMin:
    """{minute: {(direction, movement): count}} for minutes actually counted.

    Auto-detects the R/T/L/U movement row (cam2 has it on row2, cam1 on row3),
    reads approach names from the row above it (forward-filled across each
    block's blank columns), and keeps only minutes with >=1 vehicle.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    is_mv = lambda c: c is not None and str(c).strip().upper() in _MV_FROM_LETTER
    mv_idx = max(range(min(8, len(rows))),
                 key=lambda i: sum(1 for c in rows[i] if is_mv(c)))
    appr_row, mv_row = rows[mv_idx - 1], rows[mv_idx]
    col_map: dict[int, tuple[str, str]] = {}
    cur_dir = None
    for ci in range(len(mv_row)):
        name = appr_row[ci] if ci < len(appr_row) else None
        if name:
            cur_dir = str(name).strip().split()[0].upper()
        if cur_dir and is_mv(mv_row[ci]):
            col_map[ci] = (cur_dir, _MV_FROM_LETTER[str(mv_row[ci]).strip().upper()])

    out: PerMin = {}
    for r in rows[mv_idx + 1:]:
        if not isinstance(r[0], time):
            continue
        cell = {k: int(r[ci] or 0) for ci, k in col_map.items()}
        if sum(cell.values()) > 0:
            out[r[0]] = cell
    return out


def load_miovision(cam: int) -> PerMin:
    data = MIO.parse(cam)
    labels = MIO.slot_labels(data["movements"], cam)
    out: PerMin = defaultdict(lambda: defaultdict(int))
    for tm, vols in data["per_min"].items():
        t = datetime.fromisoformat(tm).time()
        for i, v in enumerate(vols):
            d = labels[i][0].strip().split()[0].upper()
            out[t][(d, labels[i][1])] += int(v)
    return out


def load_ours(cam: int) -> PerMin:
    c = sqlite3.connect(f"data/projects/{PROJECT}/project.db")
    leg_dir = {lid: _CARD_TO_DIR.get((card or "").strip().upper())
               for lid, card in c.execute(
                   "SELECT leg_id,cardinal_direction FROM legs WHERE camera_id=?", (cam,))}
    norm = {"through": "thru", "u_turn": "uturn", "left": "left", "right": "right"}
    out: PerMin = defaultdict(lambda: defaultdict(int))
    for olid, mv, ts in c.execute(
            "SELECT origin_leg_id, movement, timestamp_real FROM vehicle_events "
            "WHERE camera_id=? AND COALESCE(rejected,0)=0", (cam,)):
        d = leg_dir.get(olid)
        if ts and d:
            key = datetime.fromisoformat(ts).time().replace(second=0, microsecond=0)
            out[key][(d, norm.get(mv, mv))] += 1
    c.close()
    return out


def _agg(src: PerMin, minutes: list[time]) -> dict[tuple, int]:
    out: dict[tuple, int] = defaultdict(int)
    for m in minutes:
        for k, v in src.get(m, {}).items():
            out[k] += v
    return out


def report(cam: int) -> None:
    manual = parse_manual(MANUAL_BY_CAM[cam])
    mio, ours = load_miovision(cam), load_ours(cam)
    counted = sorted(manual.keys())
    lo, hi = counted[0].strftime("%H:%M"), counted[-1].strftime("%H:%M")
    print(f"\n########## cam{cam}  {MANUAL_BY_CAM[cam].stem} ##########")
    print(f"manual counted window: {lo}-{hi} ({len(counted)} minutes)")
    M, MI, O = _agg(manual, counted), _agg(mio, counted), _agg(ours, counted)
    print(f"{'cell':<12}{'MANUAL':>8}{'MIOVIS':>8}{'OURS':>8}"
          f"{'Mio-Man':>9}{'Ours-Man':>9}{'Ours/Man':>10}")
    tM = tMI = tO = 0
    rows_out = []
    for d in DIRECTIONS:
        for mv in MOVEMENTS:
            m, mi, o = M.get((d, mv), 0), MI.get((d, mv), 0), O.get((d, mv), 0)
            if m == 0 and mi == 0 and o == 0:
                continue
            tM += m; tMI += mi; tO += o
            ratio = f"{o/m:.2f}" if m else ("inf" if o else "-")
            print(f"{d+'-'+mv:<12}{m:>8}{mi:>8}{o:>8}{mi-m:>+9}{o-m:>+9}{ratio:>10}")
            rows_out.append((d, mv, m, mi, o))
    print(f"{'TOTAL':<12}{tM:>8}{tMI:>8}{tO:>8}{tMI-tM:>+9}{tO-tM:>+9}"
          f"{(str(round(tO/tM,2)) if tM else '-'):>10}")
    if tM:
        print(f"  net vs MANUAL:   Miovision {100*(tMI-tM)/tM:+.1f}%   OURS {100*(tO-tM)/tM:+.1f}%")
    worst = sorted(rows_out, key=lambda r: -abs(r[4] - r[2]))[:5]
    print("  largest per-movement gaps (OURS - MANUAL):")
    for d, mv, m, mi, o in worst:
        tag = "  [OURS, mio agrees w/ manual]" if abs(mi - m) <= max(3, 0.15 * m) and abs(o - m) > max(3, 0.15 * m) else ""
        print(f"    {d}-{mv:<6} manual {m:>4}  ours {o:>4} ({o-m:+d})  [miovision {mi}]{tag}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--camera", type=int, default=None,
                    help="camera id; omit to run every camera with a manual file")
    args = ap.parse_args()
    cams = [args.camera] if args.camera else sorted(MANUAL_BY_CAM)
    for cam in cams:
        report(cam)
    return 0


if __name__ == "__main__":
    sys.exit(main())
